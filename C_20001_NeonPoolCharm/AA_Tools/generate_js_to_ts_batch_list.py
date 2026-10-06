# -*- coding: utf-8 -*-
"""Generate JS-to-TS conversion batch checklist Excel."""

import os
import re
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCRIPTS = os.path.join(ROOT, "assets", "scripts")
OUTPUT = os.path.join(ROOT, "JS_to_TS_BatchList.xlsx")

SKIP_FILES = {
    "crypto-js.js", "polyglot.min.js", "zlib_min.js", "regeneratorRuntime.js",
    "debug-polyfill.js", "velocity.js", "_process.js",
    "use_v2.1.x_cc.Action.js", "use_v2.1-2.2.1_cc.Toggle_event.js",
}

PROMPT_SUFFIX = (
    "to TypeScript. Keep original .js file. Keep .js.meta file. "
    "Do NOT create .ts.meta file."
)


def rel_path(abs_path: str) -> str:
    return abs_path.replace(ROOT + os.sep, "").replace("\\", "/")


def has_ts(js_path: str) -> bool:
    return os.path.exists(os.path.splitext(js_path)[0] + ".ts")


def make_prompt(files: list[str]) -> str:
    lines = [f"Convert @{rel_path(f)}" for f in files]
    lines.append(PROMPT_SUFFIX)
    return "\n".join(lines)


def chunk_list(items: list, size: int) -> list[list]:
    return [items[i:i + size] for i in range(0, len(items), size)]


def is_obfuscated(name: str) -> bool:
    base = os.path.splitext(name)[0]
    return bool(re.match(r"^[A-Z0-9_]{5,}$", base))


def scan_js_files():
    all_files = []
    for dirpath, _, filenames in os.walk(SCRIPTS):
        for name in sorted(filenames):
            if not name.endswith(".js"):
                continue
            full = os.path.join(dirpath, name)
            all_files.append(full)
    return all_files


def build_batches():
    batches = []

    # --- Completed (reference) ---
    frame_done = sorted(
        f for f in scan_js_files()
        if "/Frame/" in f.replace("\\", "/") and has_ts(f)
    )
    if frame_done:
        batches.append((
            "✅ Frame (已完成)",
            make_prompt(frame_done),
        ))

    # --- Phase 0 ---
    newhand = os.path.join(SCRIPTS, "newHand", "newHand.js")
    if os.path.exists(newhand) and not has_ts(newhand):
        batches.append(("Phase0-newHand", make_prompt([newhand])))

    # --- Phase 1: Infrastructure (manual curated) ---
    phase1a = [
        "PlayerDataMgr.js", "PlayerDataSys.js", "ConfigDataMgr.js", "ConfigDataSys.js",
        "SystemDataMgr.js", "SystemDataSys.js", "SystemConfig.js", "ClientData.js",
        "DB.js", "defines.js",
    ]
    phase1b = [
        "PageMgr.js", "AdManager.js", "AudioManager.js", "PoolManager.js",
        "EventMgr.js", "EventDispatcher.js", "SdkHelper.js", "UrlMgr.js",
        "FileMgr.js", "UiManage.js", "frameworkManager.js", "EngineUtil.js",
    ]
    phase1c = [
        "BasePageCtrl.js", "BaseButton.js", "BaseSystem.js",
        "CocosHelper.js", "CocosHelperUtil.js",
    ]
    third = os.path.join(SCRIPTS, "third")
    for label, names in [
        ("Phase1A-DataConfig", phase1a),
        ("Phase1B-Managers", phase1b),
        ("Phase1C-BaseClasses", phase1c),
    ]:
        files = [os.path.join(third, n) for n in names if os.path.exists(os.path.join(third, n)) and not has_ts(os.path.join(third, n))]
        if files:
            batches.append((label, make_prompt(files)))

    # --- Phase 2: Huge files (one per batch) ---
    huge_order = [
        "GameConfigurations.js", "game_table.js", "List.js", "FModeConfig2.js",
        "BallLogicMgr.js", "GameServiceMgr.js", "FModeConfig1.js", "LocalServer.js",
    ]
    for i, name in enumerate(huge_order, 1):
        fp = os.path.join(third, name)
        if os.path.exists(fp) and not has_ts(fp):
            batches.append((f"Phase2-Huge-{i:02d}-{os.path.splitext(name)[0]}", make_prompt([fp])))

    # --- Phase 3: Page pairs (manual curated) ---
    phase3 = {
        "Phase3A-StartLoadingRank": [
            "StartPage.js", "StartPageCtrl.js", "LoadingPage.js", "LoadingPageCtrl.js",
            "RankPage.js", "RankPageCtrl.js", "RankPageNew.js", "MainUI.js", "MainUICtrl.js",
        ],
        "Phase3B-InGamePages": [
            "SetPageInGame.js", "SetPageInGameCtrl.js", "BuyPropPage.js", "BuyPropPageCtrl.js",
            "PropPage.js", "PropPageCtrl.js", "UsePropPage.js", "UsePropPageCtrl.js",
            "FuHuoPage.js", "FuHuoPageCtrl.js", "VideoAlertPage.js", "VideoAlertPageCtrl.js",
        ],
        "Phase3C-WithdrawShopCue": [
            "WithdrawSuccessPage.js", "WithdrawSuccessPageCtrl.js", "WithdrawItemCtrl.js",
            "CuePage.js", "CuePageCtrl.js", "MoreGamePage.js", "MoreGamePageCtrl.js",
            "DiamondToastPage.js", "DiamondToastPageCtrl.js",
        ],
    }
    for label, names in phase3.items():
        files = [os.path.join(third, n) for n in names if os.path.exists(os.path.join(third, n)) and not has_ts(os.path.join(third, n))]
        if files:
            batches.append((label, make_prompt(files)))

    # --- Phase 4: game_* ---
    phase4a = [
        "game_hall.js", "game_shop.js", "game_rank.js", "game_btn_radBall.js",
        "game_UI_alert.js", "game_UI_alert_ad.js", "game_UI_author.js",
        "game_UI_condition.js", "game_UI_reward.js", "game_UI_settting.js",
    ]
    phase4b = [
        "game_UI_radPage.js", "game_UI_tips.js", "game_table_editor.js",
        "game-table-physics-check.js", "game-table-physics-bound.js", "loadingCN.js",
    ]
    for label, names in [("Phase4A-gameUI-1", phase4a), ("Phase4B-gameUI-2", phase4b)]:
        files = [os.path.join(third, n) for n in names if os.path.exists(os.path.join(third, n)) and not has_ts(os.path.join(third, n))]
        if files:
            batches.append((label, make_prompt(files)))

    # Collect remaining third/ files not yet assigned
    assigned = set()
    for _, info in batches:
        for line in info.split("\n"):
            if line.startswith("Convert @"):
                assigned.add(line.replace("Convert @", "").strip())

    remaining = []
    for fp in scan_js_files():
        rp = rel_path(fp)
        if name := os.path.basename(fp):
            if name in SKIP_FILES:
                continue
        if "/Frame/" in rp.replace("\\", "/") and has_ts(fp):
            continue
        if has_ts(fp):
            continue
        if rp in assigned:
            continue
        if not rp.startswith("assets/scripts/third/"):
            continue
        remaining.append(fp)

    # Phase 5A: first tiny batch (curated utilities)
    phase5a_names = [
        "CoinComp.js", "FloatTipComp.js", "NodeUtil.js", "ColorUtil.js", "MathUtil.js",
        "CurrencyUtil.js", "EventHandler.js", "PageConfig.js", "PoolItem.js", "PrefabPool.js",
        "PeopleItem.js", "DrawIItem.js", "PoolNative.js", "PoolLogger.js", "NativeEventType.js",
        "AdvertEventType.js", "RequestData.js", "RequestType.js", "FormData.js", "platform.js",
    ]
    phase5a = [os.path.join(third, n) for n in phase5a_names if os.path.join(third, n) in remaining]
    if phase5a:
        batches.append(("Phase5A-Utils-Tiny", make_prompt(phase5a)))
        for f in phase5a:
            remaining.remove(f)

    # Split remaining: obfuscated vs normal
    obfuscated = sorted([f for f in remaining if is_obfuscated(os.path.basename(f))])
    normal = sorted([f for f in remaining if f not in obfuscated])

    # Phase 5: normal files by size
    def line_count(fp):
        try:
            with open(fp, "r", encoding="utf-8", errors="ignore") as fh:
                return sum(1 for _ in fh)
        except OSError:
            return 0

    tiny = [f for f in normal if line_count(f) < 80]
    small = [f for f in normal if 80 <= line_count(f) < 150]
    medium = [f for f in normal if 150 <= line_count(f) < 300]
    large = [f for f in normal if 300 <= line_count(f) < 500]

    for idx, group in enumerate(chunk_list(tiny, 20), 1):
        batches.append((f"Phase5B-Tiny-{idx:02d}", make_prompt(group)))
    for idx, group in enumerate(chunk_list(small, 12), 1):
        batches.append((f"Phase5C-Small-{idx:02d}", make_prompt(group)))
    for idx, group in enumerate(chunk_list(medium, 10), 1):
        batches.append((f"Phase5D-Medium-{idx:02d}", make_prompt(group)))
    for idx, group in enumerate(chunk_list(large, 5), 1):
        batches.append((f"Phase5E-Large-{idx:02d}", make_prompt(group)))

    # Phase 6: obfuscated ALLCAPS
    for idx, group in enumerate(chunk_list(obfuscated, 20), 1):
        batches.append((f"Phase6-Obfuscated-{idx:02d}", make_prompt(group)))

    # Skip list (reference only)
    skip_paths = []
    for name in sorted(SKIP_FILES):
        fp = os.path.join(third, name)
        if os.path.exists(fp):
            skip_paths.append(fp)
    if skip_paths:
        batches.append(("⛔ Skip-ThirdParty", "以下文件为第三方/压缩库，不建议转换：\n" + "\n".join(rel_path(f) for f in skip_paths)))

    return batches


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "BatchList"

    # Headers: col1 empty, col2 batchName, col3 batchInfo
    ws["A1"] = ""
    ws["B1"] = "batchName"
    ws["C1"] = "batchInfo"
    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(bold=True, color="FFFFFF")
    for col in ("A1", "B1", "C1"):
        ws[col].fill = header_fill
        ws[col].font = header_font
        ws[col].alignment = Alignment(horizontal="center", vertical="center")

    batches = build_batches()
    for row_idx, (name, info) in enumerate(batches, start=2):
        ws.cell(row=row_idx, column=1, value="")
        ws.cell(row=row_idx, column=2, value=name)
        cell = ws.cell(row=row_idx, column=3, value=info)
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 100
    ws.freeze_panes = "A2"

    wb.save(OUTPUT)
    print(f"Created: {OUTPUT}")
    print(f"Total batches: {len(batches)}")


if __name__ == "__main__":
    main()
